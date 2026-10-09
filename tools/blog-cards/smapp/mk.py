import os, sys, copy
from PIL import Image, ImageDraw, ImageOps, ImageFilter
from pptx import Presentation
from pptx.util import Emu
R = '/home/user/SFW-newwebsite/'
SCREEN = Image.open('smapp_shots/home.png').convert('RGB')     # real sMApp sign-in page, 2880x1800
WES = ImageOps.exif_transpose(Image.open(R + 'tools/blog-cards/photos/a-fresh-smapp.jpg')).convert('RGB')
MICRO = {'amoeba': R + 'img/sfw-amoeba-still-square.jpg',
         'scope': R + 'tools/blog-cards/photos/drive/R5A_4268.jpg'}

SUBJ = {'amoeba': (.50, .46), 'scope': (.55, .52)}   # where the subject sits in each source

def place(src, w, h, subj, tgt, z):
    """Crop src to w x h, zoomed by z, with the subject landing at tgt (fractions of the frame)."""
    a = w / h; cw = min(src.width, src.height * a) / z; ch = cw / a
    x0 = min(max(subj[0] * src.width - tgt[0] * cw, 0), src.width - cw)
    y0 = min(max(subj[1] * src.height - tgt[1] * ch, 0), src.height - ch)
    return src.crop((int(x0), int(y0), int(x0 + cw), int(y0 + ch))).resize((w, h), Image.LANCZOS)

def macbook(w):
    """A MacBook Air style laptop, front view, screen showing sMApp. Returns RGBA."""
    s = 3; W = w * s
    lid_w = int(W * .86); bez = int(lid_w * .022); sw = lid_w - 2 * bez; sh = int(sw * 10 / 16)
    lid_h = sh + 2 * bez + int(bez * .3)
    base_h = int(W * .035); H = lid_h + base_h + int(W * .03)
    im = Image.new('RGBA', (W, H), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    lx = (W - lid_w) // 2
    # shadow on the surface
    sh_ = Image.new('L', (W, H), 0); ImageDraw.Draw(sh_).ellipse((int(W*.03), lid_h + base_h - int(W*.012), int(W*.97), lid_h + base_h + int(W*.022)), fill=120)
    sh_ = sh_.filter(ImageFilter.GaussianBlur(W * .01)); im.paste((20, 15, 10, 255), (0, 0), sh_)
    # lid: aluminium edge, black glass, screen
    r = int(bez * 1.6)
    d.rounded_rectangle((lx - 3*s, 0 - 0, lx + lid_w + 3*s, lid_h + 2*s), r + 3*s, fill=(176, 178, 182, 255))
    d.rounded_rectangle((lx, 2*s, lx + lid_w, lid_h), r, fill=(12, 12, 14, 255))
    scr = SCREEN.resize((sw, sh), Image.LANCZOS)
    im.paste(scr, (lx + bez, 2*s + bez))
    # notch
    nw = int(sw * .085); d.rounded_rectangle(((W - nw)//2, 2*s + bez - 2, (W + nw)//2, 2*s + bez + int(bez*.9)), int(bez*.4), fill=(12, 12, 14, 255))
    # base
    by = lid_h
    d.polygon([(0, by + 2*s), (W, by + 2*s), (W - int(W*.012), by + base_h), (int(W*.012), by + base_h)], fill=(196, 198, 202, 255))
    d.rectangle((0, by, W, by + 2*s + 1), fill=(222, 224, 227, 255))
    d.line((int(W*.012), by + base_h - s, W - int(W*.012), by + base_h - s), fill=(140, 142, 146, 255), width=2*s)
    nw = int(W * .14); d.rounded_rectangle(((W - nw)//2, by, (W + nw)//2, by + int(base_h * .45)), int(base_h * .3), fill=(160, 162, 166, 255))
    return im.resize((w, H // s), Image.LANCZOS)

def build(src, out, key, opt):
    p = Presentation(src); sl = p.slides[0]
    SW = 1920 if 'desktop' in key else 1400; PX = p.slide_width / SW
    shapes = {s.name: s for s in sl.shapes}
    ph = shapes['Photo']; W_, H_ = round(ph.width / PX), round(ph.height / PX)
    k = 2; Wk, Hk = W_ * k, H_ * k
    micro = Image.open(MICRO[opt['micro']]).convert('RGB')
    img = Image.new('RGB', (Wk, Hk))
    bw = int(min(Wk * opt.get('wes', .5), Hk * WES.width / WES.height))
    img.paste(place(micro, Wk - bw, Hk, SUBJ[opt['micro']], opt['tgt'][key], opt.get('zoom', 1.25)), (0, 0))
    img.paste(ImageOps.fit(WES, (bw, Hk), Image.LANCZOS, centering=(.45, .45)), (Wk - bw, 0))
    tmp = f'/tmp/mk_{key}_{opt["name"]}.jpg'; img.save(tmp, quality=93)
    pic = sl.shapes.add_picture(tmp, ph.left, ph.top, ph.width, ph.height); pic.name = 'Photo'
    ph._element.addprevious(pic._element); ph._element.getparent().remove(ph._element)
    # the laptop
    paper = shapes['Paper']; px0 = paper.left / PX; py0 = paper.top / PX; pw0 = paper.width / PX
    lw = int(opt['lw'] * SW)
    lap = macbook(lw * 2); lh = lap.height / 2
    room = (py0 + 26 - 18) if not opt.get('behind') else (py0 + opt.get('sink', 150) - 18)
    if lh > room:
        lw = int(lw * room / lh); lap = macbook(lw * 2); lh = lap.height / 2
    lt = f'/tmp/mk_lap_{key}_{opt["name"]}.png'; lap.save(lt)
    x, y = opt['pos'](px0, py0, pw0, lw, lh, SW, H_)
    pic = sl.shapes.add_picture(lt, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(lw * PX)), Emu(int(lh * PX))); pic.name = 'MacBook'
    if opt.get('behind'): paper._element.addprevious(pic._element)
    else: paper._element.addnext(pic._element)
    p.save(out)

OPTS = [
 dict(name='1-amoeba-laptop-left', micro='amoeba', lw=.18, zoom=1.35, tgt={'desktop': (.74, .24), 'square': (.74, .36)},
      pos=lambda px, py, pw, lw, lh, SW, H: (px + 70, py + 24 - lh)),
 dict(name='2-microscope-laptop-left', micro='scope', lw=.18, zoom=1.15, tgt={'desktop': (.76, .30), 'square': (.74, .40)},
      pos=lambda px, py, pw, lw, lh, SW, H: (px + 70, py + 24 - lh)),
 dict(name='3-amoeba-laptop-right', micro='amoeba', lw=.18, zoom=1.35, tgt={'desktop': (.26, .24), 'square': (.28, .36)},
      pos=lambda px, py, pw, lw, lh, SW, H: (px + pw - lw - 70, py + 24 - lh)),
 dict(name='4-microscope-laptop-right', micro='scope', lw=.18, zoom=1.15, tgt={'desktop': (.24, .30), 'square': (.28, .40)},
      pos=lambda px, py, pw, lw, lh, SW, H: (px + pw - lw - 70, py + 24 - lh)),
]
if __name__ == '__main__':
    os.makedirs('mk', exist_ok=True)
    for o in OPTS:
        for key, f in [('desktop', 'sm/SFW-blog-cards--desktop-header-1920x720--brand.pptx'), ('square', 'sm/SFW-blog-cards--feature-social-1400x1400--brand.pptx')]:
            o2 = dict(o)
            if key == 'square': o2['lw'] = o['lw'] * 1.15; o2['wes'] = .5
            build(f, f'mk/sMApp-mockup-{o["name"]}--{key}.pptx', key, o2)
    print('ok')
