import os
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Emu
from mk import macbook, place, WES, R
import mk
NEM = {'nem1': (R + 'tools/blog-cards/photos/microbes/bacterial-feeding-nematode-40x-talbot-armstrong.jpg', (.42, .28)),
       'nem2': (R + 'tools/blog-cards/photos/microbes/bacterial-feeding-nematode-40x.jpg', (.70, .25))}
SCOPE = Image.open('cut/microscope-white.png')

def add_png(sl, img, x, y, w, PX, name, anchor, before):
    h = w * img.height / img.width
    p = f'/tmp/mk2_{name}_{w}.png'; img.save(p)
    pic = sl.shapes.add_picture(p, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX))); pic.name = name
    (anchor._element.addprevious if before else anchor._element.addnext)(pic._element)
    return h

def build(src, out, key, o):
    p = Presentation(src); sl = p.slides[0]
    SW = 1920 if key == 'desktop' else 1400; PX = p.slide_width / SW
    sh = {s.name: s for s in sl.shapes}; ph = sh['Photo']; paper = sh['Paper']
    W_, H_ = round(ph.width / PX), round(ph.height / PX); k = 2; Wk, Hk = W_ * k, H_ * k
    f, subj = NEM[o['bg']]; micro = Image.open(f).convert('RGB')
    img = Image.new('RGB', (Wk, Hk))
    bw = int(min(Wk * .5, Hk * WES.width / WES.height))
    img.paste(place(micro, Wk - bw, Hk, subj, o['tgt'][key], o.get('zoom', 1.2)), (0, 0))
    img.paste(ImageOps.fit(WES, (bw, Hk), Image.LANCZOS, centering=(.45, .45)), (Wk - bw, 0))
    t = f'/tmp/mk2_{key}_{o["name"]}.jpg'; img.save(t, quality=93)
    pic = sl.shapes.add_picture(t, ph.left, ph.top, ph.width, ph.height); pic.name = 'Photo'
    ph._element.addprevious(pic._element); ph._element.getparent().remove(ph._element)
    px, py, pw = paper.left / PX, paper.top / PX, paper.width / PX
    top = py + 24                                   # visible top edge of the paper
    room = top - 18
    s = o['scale'][key]
    # MacBook resting on the paper
    lw = int(SW * s['lap']); lap = macbook(lw * 2)
    if lap.height / 2 > room: lw = int(lw * room / (lap.height / 2)); lap = macbook(lw * 2)
    lx = px + s['lapx'] * pw
    add_png(sl, lap, lx, top - lap.height / 2 + 4, lw, PX, 'MacBook', paper, False)
    # microscope sticker: its cut-off base tucks behind the paper
    mw = int(SW * s['scope']); mh = mw * SCOPE.height / SCOPE.width
    tuck = mh * .12                                  # the photo's cut edge at the base hides under the paper
    if mh - tuck > room: f_ = room / (mh - tuck); mw = int(mw * f_); mh *= f_; tuck *= f_
    mx = px + s['scopex'] * pw
    add_png(sl, SCOPE, mx, top + tuck - mh, mw, PX, 'Microscope', paper, True)
    p.save(out)

OPTS = [
 dict(name='A-laptop-then-microscope', bg='nem1', zoom=1.1, tgt={'desktop': (.86, .30), 'square': (.80, .30)},
      scale={'desktop': dict(lap=.17, lapx=.06, scope=.15, scopex=.60), 'square': dict(lap=.21, lapx=.05, scope=.17, scopex=.30)}),
 dict(name='B-microscope-then-laptop', bg='nem1', zoom=1.1, tgt={'desktop': (.86, .30), 'square': (.80, .30)},
      scale={'desktop': dict(lap=.17, lapx=.30, scope=.15, scopex=.06), 'square': dict(lap=.21, lapx=.19, scope=.17, scopex=.04)}),
 dict(name='C-big-laptop-small-microscope', bg='nem1', zoom=1.1, tgt={'desktop': (.88, .28), 'square': (.80, .28)},
      scale={'desktop': dict(lap=.21, lapx=.05, scope=.11, scopex=.66), 'square': dict(lap=.26, lapx=.04, scope=.13, scopex=.34)}),
]
if __name__ == '__main__':
    os.makedirs('mk2', exist_ok=True)
    for o in OPTS:
        for key, f in [('desktop', 'sm/SFW-blog-cards--desktop-header-1920x720--brand.pptx'), ('square', 'sm/SFW-blog-cards--feature-social-1400x1400--brand.pptx')]:
            build(f, f'mk2/sMApp-option-{o["name"]}--{key}.pptx', key, o)
    print('ok')
