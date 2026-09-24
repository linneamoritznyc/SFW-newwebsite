"""Soil Health Week 2025, tablet: the band layout Linnea liked, with real
flowers. Nothing is drawn or generated: every plant is cut out of the USDA NRCS
Montana photographs in img/pd/ (knapweed buds, and bird's-foot trefoil from a
meadow), either in its own colours or as a cyanotype-style silhouette, the way
Anna Atkins printed real plants in the 1840s."""
import os, sys, math, random
OUT = sys.argv[1]
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.argv = [sys.argv[0], OUT]
import variants2 as V
B = V.B
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Emu
HERE = os.path.dirname(os.path.abspath(__file__))
W, H = V.W, V.H
PX = B.PX
BLOOMS = [2, 3, 4, 7, 8, 10, 12]
MEADOW = os.path.join(B.REPO, 'img', 'pd', 'USDA NRCS Montana (Flickr)54323527339_aeea673c88_k.jpg')

def place(s, im, name, x, y, w, rot=0, flip=False):
    if flip: im = ImageOps.mirror(im)
    h = w * im.height / im.width
    p = os.path.join(B.TMP, f'v5-{name}.png'); im.save(p)
    pic = s.shapes.add_picture(p, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    pic.name = name
    if rot: pic.rotation = rot
    return h

def silhouette(colour):
    im = Image.open(os.path.join(HERE, 'sil-knapweed.png'))
    a = im.split()[3]
    out = Image.new('RGBA', im.size, B.hexrgb(colour) + (0,)); out.putalpha(a)
    return out

def knap(): return Image.open(os.path.join(HERE, 'cut-knapweed.png'))
def bloom(i): return Image.open(os.path.join(HERE, f'cut-bloom-{BLOOMS[i % len(BLOOMS)]}.png'))

def meadow_strip(s, name, y, h):
    src = Image.open(MEADOW).convert('RGB')
    k = 2
    crop = B.cover(src.crop((0, int(src.height * .45), src.width, src.height)), W * k, int(h * k), (.5, .6), slug='none')
    mask = Image.new('L', crop.size, 0)
    from PIL import ImageDraw
    d = ImageDraw.Draw(mask)
    pts = [(i / 100 * crop.width, crop.height * .18 * (1 - math.sin(math.pi * i / 100) * .8)) for i in range(101)]
    d.polygon(pts + [(crop.width, crop.height), (0, crop.height)], fill=255)
    crop = crop.convert('RGBA'); crop.putalpha(mask)
    p = os.path.join(B.TMP, f'v5-{name}.png'); crop.save(p)
    s.shapes.add_picture(p, Emu(0), Emu(int(y * PX)), Emu(int(W * PX)), Emu(int(h * PX))).name = 'Meadow'

def base(s, c1, c2, eye, text='#FFFFFF', band_w=430, light=False):
    B.rect(s, 0, 0, W, H, '#FFFFFF', 'Page')
    if light: B.rect(s, 0, 0, band_w, H, c1, 'Ground')
    else: V.band(s, 0, 0, band_w, H, c1, c2)
    V.photo(s, f'v5-{c1}-{band_w}', band_w - 30, 0, W - band_w + 30, H, curve='left')

def words(s, eye, text='#FFFFFF', y0=150, btn_y=560):
    V.eyebrow(s, 48, y0, V.POST['cat'], eye)
    y = V.headline(s, 48, y0 + 30, 340, text)
    V.deck(s, 48, y + 20, 330, text)
    V.read_post(s, 48, btn_y, url_colour=text)

DESIGNS = []
def design(t):
    def d(fn): DESIGNS.append((t, fn)); return fn
    return d

@design('1 Cyanotype knapweed rising in the blue band')
def d1(s):
    base(s, '#3780B8', '#1F4E73', '#FFFFFF')
    place(s, silhouette('#FFFFFF'), 'sil', -40, 470, 470)
    words(s, '#FFFFFF', y0=90, btn_y=330)

@design('2 Glow cyanotype on Food Web Green')
def d2(s):
    base(s, '#156826', '#22371F', '#DBE6A7')
    place(s, silhouette('#DBE6A7'), 'sil2', 120, 420, 420, rot=-8)
    words(s, '#DBE6A7', y0=90, btn_y=330)

@design('3 Pressed trefoil along the seam, rust band')
def d3(s):
    base(s, '#9C4A24', '#7A3818', '#F4F1EA')
    random.seed(3)
    for i, (y, w, r) in enumerate([(40, 110, -20), (190, 90, 15), (330, 120, -5), (480, 95, 25), (610, 115, -15)]):
        place(s, bloom(i), f'bloom{i}', 360 + random.randint(-10, 10), y, w, rot=r, flip=i % 2)
    words(s, '#F4F1EA')

@design('4 A meadow across the foot')
def d4(s):
    base(s, '#3780B8', '#1F4E73', '#FFFFFF')
    meadow_strip(s, 'meadow4', 560, 208)
    words(s, '#FFFFFF', y0=90, btn_y=400)

@design('5 Knapweed in its own colours on Glow')
def d5(s):
    base(s, '#DBE6A7', None, '#156826', light=True)
    place(s, knap(), 'knap5', -30, 430, 460)
    words(s, '#156826', '#3C3841', y0=90, btn_y=330)

@design('6 Trefoil tucked into the corners of a plate, white page')
def d6(s):
    B.rect(s, 0, 0, W, H, '#FFFFFF', 'Page')
    V.photo(s, 'plate6', 470, 70, 500, 600, radius=8)
    for i, (x, y, w, r) in enumerate([(430, 30, 120, -25), (880, 600, 130, 20), (905, 40, 80, 10)]):
        place(s, bloom(i + 1), f'b6-{i}', x, y, w, rot=r)
    V.caption(s, 470, 684, 500, '#6A665C', V.CAPTION)
    V.eyebrow(s, 56, 170, V.POST['cat'], '#9C4A24')
    y = V.headline(s, 56, 200, 360)
    V.deck(s, 56, y + 18, 340)
    V.read_post(s, 56, 560)

@design('7 Membrane Violet with a white cyanotype, full height')
def d7(s):
    base(s, '#6E5C99', '#4C3E70', '#DBE6A7')
    im = silhouette('#FFFFFF').rotate(180, expand=True)
    place(s, im, 'sil7', 10, -20, 400)
    words(s, '#DBE6A7', y0=330, btn_y=600)

@design('8 Meadow foot and a few blooms loose above it')
def d8(s):
    base(s, '#9C4A24', '#7A3818', '#F4F1EA')
    meadow_strip(s, 'meadow8', 560, 208)
    for i, (x, y, w, r) in enumerate([(300, 470, 70, -15), (90, 520, 55, 20)]):
        place(s, bloom(i + 3), f'b8-{i}', x, y, w, rot=r)
    words(s, '#F4F1EA', y0=90, btn_y=400)

prs = Presentation(); prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
for t, fn in DESIGNS:
    sl = prs.slides.add_slide(prs.slide_layouts[6]); fn(sl); sl.notes_slide.notes_text_frame.text = t
out = os.path.join(OUT, 'Tablet-headers-soil-health-week-flowers.pptx'); prs.save(out); print(out)
