import sys, os
S = sys.argv[1]
sys.argv = [sys.argv[0], os.path.join(S, 'lap')]
sys.path.insert(0, S)
import smapp_laptop as L
from pptx import Presentation
M = {'desktop-header-1920x720': 'Desktop-1920x720', 'tablet-header-1024x768': 'Tablet-1024x768', 'thumbnail-1200x800': 'Thumbnail-3x2-1200x800',
     'feature-social-1400x1400': 'Social-Square-1400x1400', 'mobile-header-750x1000': 'Mobile-750x1000'}
for key, name in M.items():
    f = os.path.join(S, 'FINAL3', f'SFW-blog-headers-{name}.pptx')
    prs = Presentation(f); sl = prs.slides._sldIdLst
    W, H, BIG = L.SIZES[key]
    L.v_split(prs.slides.add_slide(prs.slide_layouts[6]), key, W, H, BIG)
    first = sl[0]; prs.part.drop_rel(first.rId); sl.remove(first)
    new = sl[-1]; sl.remove(new); sl.insert(0, new)
    prs.save(f); print(name, len(prs.slides))
