"""Soil Health Week 2025, tablet: flowers as real photographs, not cut-outs.
The meadow (bird's-foot trefoil) and knapweed photographs from img/pd/ are used
whole, as panels, tinted into brand colours or left natural, beside the event
photograph. Linnea's black READ POST box on every one."""
import os, sys, math
OUT = sys.argv[1]
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.argv = [sys.argv[0], OUT]
import variants2 as V
B = V.B
from PIL import Image, ImageOps, ImageEnhance, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN
W, H, PX = V.W, V.H, B.PX
PD = os.path.join(B.REPO, 'img', 'pd')
MEADOW = os.path.join(PD, 'USDA NRCS Montana (Flickr)54323527339_aeea673c88_k.jpg')
KNAP = os.path.join(PD, 'USDA NRCS Montana (Flickr)54321702858_5440c272f6_k.jpg')
GREEN, MOSS, BLUE, BLUED, RUST, RUSTD, VIOLET, VIOD, GLOW, SOIL = '#156826', '#22371F', '#3780B8', '#1F4E73', '#9C4A24', '#7A3818', '#6E5C99', '#4C3E70', '#DBE6A7', '#4F3433'

def flora(src, w, h, focus=(.5, .5), tone=None, dark=.0, curve=None, radius=0):
    im = Image.open(src).convert('RGB')
    k = max(1, min(2, min(im.width / w, im.height / h)))
    im = ImageOps.fit(im, (int(w * k), int(h * k)), Image.LANCZOS, centering=focus)
    if tone:
        g = ImageOps.autocontrast(ImageOps.grayscale(im), cutoff=1)
        im = ImageOps.colorize(g, B.hexrgb(tone[0]), B.hexrgb(tone[1]), mid=B.hexrgb(tone[2]) if len(tone) > 2 else None)
    if dark:
        im = ImageEnhance.Brightness(im).enhance(1 - dark)
    if curve or radius:
        mask = Image.new('L', im.size, 0); d = ImageDraw.Draw(mask)
        if radius: d.rounded_rectangle((0, 0, im.width - 1, im.height - 1), radius=int(radius * k), fill=255)
        else:
            amp = im.height * .09
            pts = [(i / 100 * im.width, amp * (.55 - .45 * i / 100 + .35 * math.sin(math.pi * i / 100))) for i in range(101)]
            d.polygon(pts + [(im.width, im.height), (0, im.height)], fill=255)
        im = im.convert('RGBA'); im.putalpha(mask)
    return im

def put(s, im, name, x, y, w, h):
    ext = 'png' if im.mode == 'RGBA' else 'jpg'
    p = os.path.join(B.TMP, f'v6-{name}.{ext}')
    im.save(p, quality=93) if ext == 'jpg' else im.save(p)
    s.shapes.add_picture(p, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX))).name = name

def words(s, x, y0, w, text='#FFFFFF', eye='#FFFFFF', btn_y=None):
    V.eyebrow(s, x, y0, V.POST['cat'], eye)
    y = V.headline(s, x, y0 + 30, w, text)
    y = V.deck(s, x, y + 18, w * .95, text)
    V.read_post(s, x, btn_y or y + 40, url_colour=text)

def event(s, name, x=400, curve='left'):
    V.photo(s, 'v6ev-' + name, x, 0, W - x, H, curve=curve)

DESIGNS = []
def design(t):
    def d(fn): DESIGNS.append((t, fn)); return fn
    return d

@design('1 The meadow as the band, toned Food Web Green')
def d1(s):
    V.band(s, 0, 0, 430, H, GREEN, MOSS)
    put(s, flora(MEADOW, 430, 340, (.45, .6), tone=(MOSS, GLOW, GREEN), curve=True), 'meadow1', 0, 428, 430, 340)
    event(s, '1')
    words(s, 48, 70, 340, eye=GLOW, btn_y=330)

@design('2 Knapweed as the band, its own colours')
def d2(s):
    put(s, flora(KNAP, 430, H, (.42, .4), dark=.35), 'knap2', 0, 0, 430, H)
    event(s, '2')
    words(s, 48, 400, 340, eye=GLOW, btn_y=610)

@design('3 Three panels: blue, knapweed, the event')
def d3(s):
    V.band(s, 0, 0, 380, H, BLUE, BLUED)
    put(s, flora(KNAP, 200, H, (.48, .35)), 'knap3', 380, 0, 200, H)
    V.photo(s, 'v6ev3', 580, 0, W - 580, H)
    words(s, 44, 150, 300, btn_y=560)

@design('4 The website hero, with a second print of the meadow')
def d4(s):
    B.rect(s, 0, 0, W, H, '#FFFFFF', 'Page')
    V.photo(s, 'v6ev4', 440, 60, 540, 540, radius=8)
    put(s, flora(MEADOW, 230, 230, (.3, .55), radius=8), 'meadow4', 370, 470, 230, 230)
    V.eyebrow(s, 56, 150, V.POST['cat'], GREEN)
    y = V.headline(s, 56, 180, 300, SOIL)
    V.deck(s, 56, y + 18, 280, SOIL)
    V.read_post(s, 56, 560)

@design('5 Rust band with the meadow printed into it')
def d5(s):
    V.band(s, 0, 0, 430, H, RUST, RUSTD)
    put(s, flora(MEADOW, 430, 340, (.55, .55), tone=(RUSTD, '#F1C9A5', RUST), curve=True), 'meadow5', 0, 428, 430, 340)
    event(s, '5')
    words(s, 48, 70, 340, eye='#F4F1EA', btn_y=330)

@design('6 Event above, meadow curving up into a green foot')
def d6(s):
    V.band(s, 0, 0, W, H, GREEN, MOSS, 'Green')
    V.photo(s, 'v6ev6', 0, 0, W, 430, curve='bottom')
    put(s, flora(MEADOW, 300, 338, (.5, .6)), 'meadow6', W - 300, 430, 300, 338)
    words(s, 56, 454, 600, eye=GLOW)

@design('7 Knapweed field full bleed, the event as a print on it')
def d7(s):
    put(s, flora(KNAP, W, H, (.5, .45), dark=.3), 'knap7', 0, 0, W, H)
    V.photo(s, 'v6ev7', 470, 70, 500, 600, radius=8)
    V.caption(s, 470, 684, 500, '#F4F1EA', V.CAPTION)
    words(s, 56, 170, 360, eye=GLOW, btn_y=560)

@design('8 Violet band, the meadow toned violet into glow')
def d8(s):
    V.band(s, 0, 0, 430, H, VIOLET, VIOD)
    put(s, flora(MEADOW, 430, 340, (.4, .6), tone=(VIOD, GLOW, VIOLET), curve=True), 'meadow8', 0, 428, 430, 340)
    event(s, '8')
    words(s, 48, 70, 340, eye=GLOW, btn_y=330)

@design('9 Blue band, knapweed toned blue like a cyanotype')
def d9(s):
    V.band(s, 0, 0, 430, H, BLUE, BLUED)
    put(s, flora(KNAP, 430, 400, (.4, .3), tone=('#10304A', '#F4F1EA', BLUE), curve=True), 'knap9', 0, 368, 430, 400)
    event(s, '9')
    words(s, 48, 60, 340, btn_y=300)

@design('10 White page, two prints side by side, flowers and people')
def d10(s):
    B.rect(s, 0, 0, W, H, '#FFFFFF', 'Page')
    put(s, flora(MEADOW, 300, 440, (.3, .6), radius=8), 'meadow10', 56, 56, 300, 440)
    V.photo(s, 'v6ev10', 380, 56, 588, 440, radius=8)
    V.eyebrow(s, 56, 530, V.POST['cat'], GREEN)
    V.headline(s, 56, 558, 600, SOIL)
    V.deck(s, 56, 620, 560, SOIL)
    V.read_post(s, 780, 560)

prs = Presentation(); prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
for t, fn in DESIGNS:
    sl = prs.slides.add_slide(prs.slide_layouts[6]); fn(sl); sl.notes_slide.notes_text_frame.text = t
out = os.path.join(OUT, 'Tablet-headers-soil-health-week-flower-photos.pptx'); prs.save(out); print(out)
