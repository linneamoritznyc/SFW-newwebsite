import sys, os
S = sys.argv[1]
sys.argv = [sys.argv[0], os.path.join(S, 'promo')]
sys.path.insert(0, S)
import smapp_promo as P
from pptx import Presentation
M = {'desktop-header-1920x720': 'Desktop-1920x720', 'tablet-header-1024x768': 'Tablet-1024x768', 'thumbnail-1200x800': 'Thumbnail-3x2-1200x800',
     'feature-social-1400x1400': 'Social-Square-1400x1400', 'mobile-header-750x1000': 'Mobile-750x1000'}
for key, name in M.items():
    f = os.path.join(S, 'FINAL2', f'SFW-blog-headers-{name}.pptx')
    prs = Presentation(f)
    # drop the old sMApp slide (the first one, stand-in photo)
    sl = prs.slides._sldIdLst
    P.slide_for(prs, key)                     # add first, so the new part name is unique
    first = sl[0]; prs.part.drop_rel(first.rId); sl.remove(first)
    new = sl[-1]; sl.remove(new); sl.insert(0, new)
    prs.save(f); print(name, len(prs.slides))
