"""sMApp launch header: three iPhones showing the app, on Food Web Green with a
dot grid, Linnea's torn-paper card at the bottom left. The phone screens are
redrawn from Linnea's screenshots of soilmapp.com (portfolio, assessments,
clients); each screen is its own picture so it can be swapped for the real
screenshot in Canva with Replace."""
import os, sys
OUT = sys.argv[1]; os.makedirs(OUT, exist_ok=True)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.argv = [sys.argv[0], OUT]
import build as B
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
PX = B.PX
TMP = os.path.join(OUT, '_promo'); os.makedirs(TMP, exist_ok=True)
F = os.path.expanduser('~/.fonts')
def f(name, s): return ImageFont.truetype(os.path.join(F, name), s)
SANS, SANSB = 'SourceSans3-Regular.ttf', 'Montserrat-Bold.ttf'
BLUE, DGREEN, GOLD, GREENBTN = (46, 123, 183), (34, 72, 58), (205, 154, 46), (76, 175, 80)
LOGO = Image.open(os.path.join(B.REPO, 'img', 'sfwlogo-240.png')).convert('RGBA')
WLOGO = Image.new('RGBA', LOGO.size, (255, 255, 255, 0)); WLOGO.putalpha(LOGO.split()[3])

def shell(w=1024, h=2216):
    """Phone screen: blue app bar with the white logo, the left icon rail."""
    im = Image.new('RGB', (w, h), (250, 250, 250)); d = ImageDraw.Draw(im)
    d.rectangle((0, 0, w, 120), fill=(255, 255, 255))
    d.text((70, 38), '1:22', font=f(SANSB, 44), fill=(20, 20, 20))
    d.rectangle((0, 120, w, 290), fill=BLUE)
    for y in (180, 205, 230): d.rectangle((50, y, 120, y + 8), fill='white')
    lg = WLOGO.resize((150, 150)); im.paste(lg, (165, 130), lg)
    d.ellipse((w - 160, 150, w - 50, 260), fill='white')
    d.ellipse((w - 122, 172, w - 88, 206), fill=BLUE); d.pieslice((w - 135, 205, w - 75, 260), 180, 360, fill=BLUE)
    d.line((145, 290, 145, h), fill=(225, 225, 225), width=3)
    for i in range(7):
        y = 360 + i * 125
        d.rounded_rectangle((45, y, 105, y + 50), 10, fill=(110, 110, 110))
    return im, d

def portfolio():
    im, d = shell()
    d.rounded_rectangle((190, 320, 980, 1060), 50, fill=DGREEN)
    d.ellipse((620, 230, 1150, 760), fill=(52, 96, 78))
    d.rounded_rectangle((240, 380, 590, 450), 35, fill=(70, 110, 92))
    d.text((265, 392), 'Consultant portfolio', font=f(SANS, 38), fill='white')
    d.text((240, 490), 'Your client', font=f(SANSB, 80), fill='white')
    d.text((240, 585), 'portfolio', font=f(SANSB, 80), fill='white')
    for i, t in enumerate(['See where every client stands, what', 'changed recently, and where your', 'attention will have the most value.']):
        d.text((240, 700 + i * 54), t, font=f(SANS, 40), fill=(225, 235, 230))
    d.rounded_rectangle((240, 900, 640, 1000), 12, fill=GOLD)
    d.text((290, 925), '+  ADD CLIENT', font=f(SANSB, 40), fill=(40, 30, 10))
    for i, (n, lab) in enumerate([('3', 'Clients'), ('5', 'Projects'), ('3', 'Locations'), ('36', 'Assessments')]):
        y = 1110 + i * 255
        d.rounded_rectangle((190, y, 980, y + 220), 30, outline=(225, 225, 225), width=3, fill='white')
        d.ellipse((240, y + 50, 360, y + 170), fill=(230, 238, 234))
        d.rounded_rectangle((275, y + 85, 325, y + 135), 6, fill=DGREEN)
        d.text((400, y + 38), n, font=f(SANSB, 70), fill=(25, 25, 25))
        d.text((400, y + 130), lab, font=f(SANS, 44), fill=(90, 90, 90))
    return im

def table(title, cols, rows, footer=None):
    im, d = shell()
    d.text((190, 340), title, font=f(SANS, 72), fill=(30, 20, 20))
    d.rounded_rectangle((800, 330, 980, 430), 10, fill=GREENBTN)
    d.text((838, 355), 'ADD', font=f(SANSB, 44), fill='white')
    top = 470
    if title == 'Assessments':
        d.rounded_rectangle((190, top, 980, top + 120), 8, outline=(215, 215, 215), width=3, fill='white')
        d.text((225, top + 35), 'Filters', font=f(SANS, 44), fill=(40, 40, 40)); top += 170
    d.rounded_rectangle((190, top, 980, 2100), 12, outline=(220, 220, 220), width=3, fill='white')
    for i, x in enumerate((600, 700, 800, 900)):
        d.rounded_rectangle((x, top + 50, x + 40, top + 90), 6, fill=(120, 120, 120))
    y = top + 160
    d.line((190, y - 20, 980, y - 20), fill=(225, 225, 225), width=3)
    for i, c in enumerate(cols): d.text((230 + i * 260, y + 30), c, font=f(SANSB, 38), fill=(30, 30, 30))
    y += 130
    for r in rows:
        d.line((190, y, 980, y), fill=(225, 225, 225), width=3)
        for i, c in enumerate(r): d.text((230 + i * 260, y + 50), c, font=f(SANS, 40), fill=(40, 40, 40))
        y += 150
    if footer:
        d.line((190, y, 980, y), fill=(225, 225, 225), width=3)
        d.text((480, y + 50), footer, font=f(SANS, 40), fill=(60, 60, 60))
    return im

def phone(screen, name):
    """An iPhone: black body, rounded screen, dynamic island, soft shadow."""
    sw, sh = screen.size
    bez = 34
    W, H = sw + 2 * bez, sh + 2 * bez
    pad = 80
    out = Image.new('RGBA', (W + 2 * pad, H + 2 * pad), (0, 0, 0, 0))
    sm = Image.new('L', out.size, 0); ImageDraw.Draw(sm).rounded_rectangle((pad, pad + 30, pad + W, pad + H + 30), 150, fill=120)
    out.putalpha(0); shadow = Image.new('RGBA', out.size, (10, 25, 15, 0)); shadow.putalpha(sm.filter(ImageFilter.GaussianBlur(40)))
    out = Image.alpha_composite(out, shadow)
    body = Image.new('RGBA', out.size, (0, 0, 0, 0)); bd = ImageDraw.Draw(body)
    bd.rounded_rectangle((pad, pad, pad + W, pad + H), 150, fill=(22, 22, 24))
    bd.rounded_rectangle((pad + 4, pad + 4, pad + W - 4, pad + H - 4), 146, outline=(70, 70, 74), width=4)
    out = Image.alpha_composite(out, body)
    m = Image.new('L', screen.size, 0); ImageDraw.Draw(m).rounded_rectangle((0, 0, sw, sh), 118, fill=255)
    scr = screen.convert('RGBA'); scr.putalpha(m)
    out.alpha_composite(scr, (pad + bez, pad + bez))
    ImageDraw.Draw(out).rounded_rectangle((pad + W // 2 - 150, pad + bez + 28, pad + W // 2 + 150, pad + bez + 108), 40, fill=(10, 10, 10))
    p = os.path.join(TMP, f'phone-{name}.png'); out.save(p)
    return p, out.size

SCREENS = [
 ('portfolio', portfolio()),
 ('assessments', table('Assessments', ['Assessment', 'Location'], [['CLP Scena...', ''], ['CLP Scena...', ''], ['Pasture 1', ''], ['workshop ...', ''], ['Mentor tea', ''], ['NM pre-wo...', ''], ['compost f...', '']])),
 ('clients', table('Clients', ['Client', 'Organi...', 'Country'], [['Awesome ...', 'hiz biz', 'usa'], ['Giant Pum...', 'over there', 'USA'], ["Wes's hom...", 'Foothill Bio...', 'usa']], '1-3 of 3')),
]
PHONES = [phone(s, n) for n, s in SCREENS]

def dots(w, h, step=28):
    im = Image.new('RGBA', (w, h), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    for y in range(step // 2, h, step):
        for x in range(step // 2, w, step): d.ellipse((x - 2, y - 2, x + 2, y + 2), fill=(255, 255, 255, 34))
    p = os.path.join(TMP, f'dots-{w}x{h}.png'); im.save(p); return p

POST = dict(slug='a-fresh-smapp', cat='Microscopy', head='A Fresh sMApp',
            deck='After lots of suggestions and lots of effort, our renovated SFW Microscopy App feels brand new again')

# per size: the phones' area (x, y, height of the middle phone) and the paper (x, y, w)
LAYOUT = {
 'desktop-header-1920x720':  dict(W=1920, H=720,  phones=(1000, 30, 640), card=(56, 0, 860), big=56),
 'feature-social-1400x1400': dict(W=1400, H=1400, phones=(160, 40, 860), card=(64, 0, 1272), big=78),
 'tablet-header-1024x768':   dict(W=1024, H=768,  phones=(520, 30, 900), card=(36, 0, 560), big=44),
 'thumbnail-1200x800':       dict(W=1200, H=800,  phones=(600, 30, 900), card=(40, 0, 620), big=46),
 'mobile-header-750x1000':   dict(W=750,  H=1000, phones=(90, 30, 560),  card=(26, 0, 698), big=46),
}

def slide_for(prs, key):
    L = LAYOUT[key]; W, H = L['W'], L['H']
    s = prs.slides.add_slide(prs.slide_layouts[6])
    B.V_band = None
    sh = s.shapes.add_shape(1, 0, 0, Emu(W * PX), Emu(H * PX)); sh.name = 'Green'
    sh.line.fill.background(); f_ = sh.fill; f_.gradient(); f_.gradient_angle = 68
    st = f_.gradient_stops; st[0].color.rgb = B.RGBColor(0x15, 0x68, 0x26); st[0].position = 0; st[1].color.rgb = B.RGBColor(0x22, 0x37, 0x1F); st[1].position = 1
    s.shapes.add_picture(dots(W, H), 0, 0, Emu(W * PX), Emu(H * PX)).name = 'Dot grid'
    # three phones: middle one in front and bigger, the others behind and a little lower
    x0, y0, mh = L['phones']
    (p, (pw, ph)) = PHONES[0]
    scale = mh / ph
    mw = pw * scale
    side = .86
    order = [(1, x0 - mw * .02, y0 + mh * .1, side), (2, x0 + mw * 1.02, y0 + mh * .1, side), (0, x0 + mw * .5 - mw * .5 + mw * .5, y0, 1.0)]
    if W < H or key.startswith('feature'):
        cx = W / 2
        order = [(1, cx - mw * 1.25, y0 + mh * .08, side), (2, cx + mw * .4, y0 + mh * .08, side), (0, cx - mw / 2, y0, 1.0)]
    else:
        # fit the three phones between x0 and the right edge
        avail = W - x0 - 24
        if mw * (1.15 + side) > avail:
            k = avail / (mw * (1.15 + side)); scale *= k; mw *= k
        order = [(1, x0, y0 + mh * .08, side), (2, x0 + mw * 1.15, y0 + mh * .08, side), (0, x0 + mw * .55, y0, 1.0)]
    for i, x, y, sc in order:
        path, (w_, h_) = PHONES[i]
        s.shapes.add_picture(path, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w_ * scale * sc * PX)), Emu(int(h_ * scale * sc * PX))).name = f'Phone {SCREENS[i][0]}'
    # the paper at the bottom
    cx_, _, cw = L['card']; BIG = L['big']; SMALL = round(BIG * .52); pad = BIG * .8; tw = cw - 2 * pad
    hl = B.wrap(POST['head'], B.font(B.F_HEAD, BIG), tw * .95)
    dl = B.wrap(POST['deck'], B.font(B.F_HEAD, SMALL), tw * .9)
    btn_h = SMALL * 2.0
    ch = int(pad + BIG * 1.08 + BIG * len(hl) + SMALL * .6 + SMALL * 1.25 * len(dl) + SMALL * 1.3 + btn_h + SMALL * 1.3 + pad * .8)
    cy = H - ch - cx_
    png = B.torn_paper(key, 'smapp-promo', cw, ch); m = B.PAPER_MARGIN
    s.shapes.add_picture(png, Emu(int((cx_ - m) * PX)), Emu(int((cy - m) * PX)), Emu(int((cw + 2 * m) * PX)), Emu(int((ch + 2 * m) * PX))).name = 'Paper'
    y = cy + pad
    B.text(s, cx_ + pad, y, tw, BIG * 1.1, POST['cat'], 'Montserrat', BIG, B.INK, 'Category', bold=True, spacing=-BIG * .75 * 3, line=BIG); y += BIG * 1.08
    B.text(s, cx_ + pad, y, tw, BIG * len(hl) + 4, '\v'.join(hl), 'Source Sans 3', BIG, B.INK, 'Headline', spacing=-BIG * .75 * 4, line=BIG); y += BIG * len(hl) + SMALL * .6
    B.text(s, cx_ + pad + 2, y, tw * .9, SMALL * 1.25 * len(dl) + 4, '\v'.join(dl), 'Source Sans 3', SMALL, B.INK, 'Deck', line=SMALL * 1.25); y += SMALL * 1.25 * len(dl)
    by = y + SMALL * 1.3
    bw = B.font(B.F_BTN, SMALL * 1.1).getlength('READ POST') + SMALL * 2.2
    B.rect(s, cx_ + pad, by, bw, btn_h, B.INK, 'Button')
    B.text(s, cx_ + pad, by, bw, btn_h, 'READ POST', 'EB Garamond', SMALL * 1.1, '#FFFFFF', 'Button label', align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=SMALL * .75 * 2)
    B.cursor(s, cx_ + pad + bw - SMALL * .5, by + btn_h * .55, SMALL * 1.5)
    B.text(s, cx_ + pad, by + btn_h + SMALL * .3, tw, SMALL, 'soilmapp.com  ·  www.soilfoodweb.com', 'Source Sans 3', SMALL * .78, B.INK, 'Web address', line=SMALL)
    # the logo in white, top left
    lg = os.path.join(TMP, 'wlogo.png'); WLOGO.save(lg)
    ls = BIG * 2.2
    s.shapes.add_picture(lg, Emu(int(cx_ * PX)), Emu(int(cx_ * .8 * PX)), Emu(int(ls * PX)), Emu(int(ls * PX))).name = 'Logo'

for key, L in LAYOUT.items():
    prs = Presentation(); prs.slide_width, prs.slide_height = Emu(L['W'] * PX), Emu(L['H'] * PX)
    slide_for(prs, key)
    out = os.path.join(OUT, f'smapp-launch--{key}.pptx'); prs.save(out); print(out)
