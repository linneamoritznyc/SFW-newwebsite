"""Advanced Programs Are Reopening, desktop 1920x720: the brochure band layouts
Linnea liked, pushed further, across the brand palette."""
import os, sys
OUT = sys.argv[1]
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.argv = [sys.argv[0], OUT]
import variants2 as V
B = V.B
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Emu

POST = next(p for p in B.POSTS if p['slug'] == 'advanced-programs-reopening')
B._CURRENT['slug'] = POST['slug']
V.POST = POST
V.SRC = ImageOps.exif_transpose(Image.open(B.photo_for(POST))).convert('RGB')
V.CAPTION = 'Students and mentors practice microscopy together at our workshop in Costa Rica, March 2025.'
W, H = 1920, 720
V.W, V.H = W, H
V.HEAD, V.DECK, V.EYE, V.CAP = 60, 24, 16, 17
PAIRS = {
 'green': ('#156826', '#22371F', '#DBE6A7'), 'blue': ('#3780B8', '#1F4E73', '#FFFFFF'),
 'violet': ('#6E5C99', '#4C3E70', '#DBE6A7'), 'rust': ('#9C4A24', '#7A3818', '#F4F1EA'),
 'soil': ('#4F3433', '#231F1D', '#C89B7B'), 'scope': ('#3C3841', '#231F1D', '#9E8FC2'),
 'olive': ('#5F6A3C', '#22371F', '#DBE6A7'),
}

def left(s, c):
    c1, c2, eye = PAIRS[c]
    B.rect(s, 0, 0, W, H, '#FFFFFF', 'Page')
    V.band(s, 0, 0, 800, H, c1, c2)
    V.photo(s, 'l-' + c, 680, 0, 1240, H, curve='left')
    V.eyebrow(s, 72, 170, POST['cat'], eye)
    V.headline(s, 72, 204, 560, '#FFFFFF')
    V.read_post(s, 72, 500, url_colour='#FFFFFF')
    V.caption(s, 72, 640, 560, '#FFFFFF', V.CAPTION)

def right(s, c):
    c1, c2, eye = PAIRS[c]
    B.rect(s, 0, 0, W, H, '#FFFFFF', 'Page')
    V.band(s, 1200, 0, 720, H, c2, c1)
    V.photo(s, 'r-' + c, 0, 0, 1240, H)
    V.eyebrow(s, 1290, 170, POST['cat'], eye)
    V.headline(s, 1290, 204, 560, '#FFFFFF')
    V.read_post(s, 1290, 500, url_colour='#FFFFFF')
    V.caption(s, 1290, 640, 560, '#FFFFFF', V.CAPTION)

def framed(s, c):
    # colour field with the photo as a plate inside it, curve on its lower edge
    c1, c2, eye = PAIRS[c]
    V.band(s, 0, 0, W, H, c1, c2, 'Colour')
    V.photo(s, 'f-' + c, 820, 60, 1040, 560, curve='bottom')
    V.caption(s, 820, 636, 1040, '#FFFFFF', V.CAPTION)
    V.eyebrow(s, 72, 170, POST['cat'], eye)
    V.headline(s, 72, 204, 640, '#FFFFFF')
    V.read_post(s, 72, 500, url_colour='#FFFFFF')

def light(s, ground, text, eye, name):
    B.rect(s, 0, 0, W, H, ground, 'Ground')
    V.photo(s, 'g-' + name, 680, 0, 1240, H, curve='left')
    V.eyebrow(s, 72, 170, POST['cat'], eye)
    V.headline(s, 72, 204, 560, text)
    V.read_post(s, 72, 500, url_colour=text)
    V.caption(s, 72, 640, 560, text, V.CAPTION)

def panel(s, c):
    # the brochure panel exactly: colour header with eyebrow and headline, photo under it, curve, caption on white
    c1, c2, eye = PAIRS[c]
    B.rect(s, 0, 0, W, H, '#FFFFFF', 'Page')
    V.band(s, 0, 0, W, 210, c1, c2)
    V.eyebrow(s, 72, 58, POST['cat'], eye)
    V.headline(s, 72, 90, 1400, '#FFFFFF')
    V.photo(s, 'p-' + c, 0, 210, W, 420, curve='bottom')
    V.caption(s, 72, 650, 1100, '#6A665C', V.CAPTION)
    V.read_post(s, 1560, 630)

def specimen(s, c):
    # the band carries a real organism from the Foundation's microscopy, floating
    c1, c2, eye = PAIRS[c]
    B.rect(s, 0, 0, W, H, '#FFFFFF', 'Page')
    V.band(s, 0, 0, 800, H, c1, c2)
    V.photo(s, 's-' + c, 680, 0, 1240, H, curve='left')
    cut = os.path.join(B.REPO, 'img', 'uploads', '2.png')
    im = Image.open(cut); cw = 170; chh = int(im.height * cw / im.width)
    s.shapes.add_picture(cut, Emu(int(470 * B.PX)), Emu(int(60 * B.PX)), Emu(int(cw * B.PX)), Emu(int(chh * B.PX))).name = 'Specimen'
    V.eyebrow(s, 72, 250, POST['cat'], eye)
    V.headline(s, 72, 284, 560, '#FFFFFF')
    V.read_post(s, 72, 540, url_colour='#FFFFFF')
    V.caption(s, 72, 650, 560, '#FFFFFF', V.CAPTION)

def triptych(s, a, b, c):
    # the brochure's three panels side by side, as three colour columns
    B.rect(s, 0, 0, W, H, '#FFFFFF', 'Page')
    for i, k in enumerate((a, b, c)):
        c1, c2, _ = PAIRS[k]; V.band(s, i * 36, 0, 36, H, c1, c2, f'Column {i+1}')
    V.photo(s, 't', 1000, 0, 920, H, curve='left')
    eye = PAIRS[a][0]
    V.eyebrow(s, 180, 190, POST['cat'], eye)
    V.headline(s, 180, 224, 720, '#4F3433')
    V.read_post(s, 180, 500)
    V.caption(s, 180, 640, 700, '#6A665C', V.CAPTION)

SLIDES = [(panel, ('blue',)), (panel, ('rust',)), (specimen, ('violet',)), (specimen, ('blue',)), (triptych, ('blue', 'green', 'violet')),
          (left, ('blue',)), (left, ('green',)), (left, ('violet',)), (left, ('rust',)), (left, ('soil',)), (left, ('olive',)),
          (right, ('blue',)), (right, ('scope',)),
          (framed, ('blue',)), (framed, ('violet',)),
          (light, ('#DBE6A7', '#3C3841', '#156826', 'glow')), (light, ('#9E8FC2', '#231F1D', '#231F1D', 'membrane'))]
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(W * B.PX), Emu(H * B.PX)
for fn, a in SLIDES:
    sl = prs.slides.add_slide(prs.slide_layouts[6]); fn(sl, *a)
    sl.notes_slide.notes_text_frame.text = f'{fn.__name__}: {a[-1]}'
out = os.path.join(OUT, 'Desktop-headers-advanced-programs-versions.pptx'); prs.save(out); print(out)
