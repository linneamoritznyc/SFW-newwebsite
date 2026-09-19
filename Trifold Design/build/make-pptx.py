"""Build an editable PowerPoint of the two trifold spreads.

    python3 "Trifold Design/build/make-pptx.py"

Canva imports a PDF by guessing where the text boxes are, because a PDF holds
only glyphs at coordinates. It guesses wrong on tight layouts: words lose their
spaces and blocks reorder. A PPTX holds real text boxes, so nothing is guessed.

Run extract-layout.js first: it writes pptx/layout.json and the two text-free
background renders this reads.
"""
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "pptx"
OUT = ROOT / "pptx" / "SFW-brochure-2026-EDITABLE.pptx"

PX_TO_PT = 0.75          # CSS px at 96 dpi -> points

# The artwork is 11.25 x 8.75in: the 11 x 8.5in trim plus 0.125in of bleed all
# round. Canva resizes any slide that is not a size it recognises, and scaling
# the shapes without scaling the type re-wraps every box. So the deck is built
# at trim size, with the bleed cropped off the background and every coordinate
# moved in by the same 0.125in. Bleed is no use in Canva anyway; it matters
# only to the press file.
BLEED = 0.125
SLACK = 0.09
TRIM_W, TRIM_H = 11.0, 8.5
# Baked line breaks mean each line is its own paragraph, so a justified
# paragraph would stretch every line to the full box width. Left is what the
# browser's last line of a justified block already looks like.
ALIGN = {"start": PP_ALIGN.LEFT, "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
         "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.LEFT}

# PowerPoint hangs the first line from the top of the box by its line height,
# where CSS centres the glyphs in the line box. Shifting each box up by the
# half-leading puts the first baseline where the browser put it. Verified by
# rendering the deck back out and diffing against the browser render.
HALF_LEADING = True


def parse_color(css):
    m = re.match(r"rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?\)", css)
    if not m:
        return RGBColor(0, 0, 0), 1.0
    r, g, b = (int(round(float(m.group(i)))) for i in (1, 2, 3))
    a = float(m.group(4)) if m.group(4) is not None else 1.0
    return RGBColor(r, g, b), a


def set_alpha(run, alpha):
    """DrawingML carries opacity as a child of the run's solid fill."""
    if alpha >= 0.999:
        return
    fill = run.font.color._xFill  # solidFill, already present from set colour
    clr = fill.find(qn("a:srgbClr"))
    if clr is None:
        return
    el = clr.makeelement(qn("a:alpha"), {"val": str(int(round(alpha * 100000)))})
    clr.append(el)


def set_spacing(run, px):
    """Letter-spacing, in hundredths of a point, on the run properties."""
    if abs(px) < 0.01:
        return
    run.font._rPr.set("spc", str(int(round(px * PX_TO_PT * 100))))


def collapse(text):
    """HTML whitespace collapsing: the source is indented, the render is not."""
    return re.sub(r"\s+", " ", text)


def add_block(slide, b):
    lines = b.get("lines") or []
    left_aligned = b["align"] in ("start", "left", "justify")
    # Left-aligned text starts where the browser actually drew it, which is not
    # the element's edge when a heading sits beside an icon. Centred and
    # right-aligned text keeps the element box so it stays centred in it.
    x = lines[0]["left"] if (lines and left_aligned) else b["x"]
    # Canva ignores wrap="none" and re-wraps to the box width, so each box has
    # to be at least as wide as its widest measured line. SLACK covers the
    # difference between the browser's text metrics and another engine's.
    line_w = max((ln["right"] for ln in lines), default=b["x"] + b["w"]) - x
    w = max(line_w + SLACK, b["x"] + b["w"] - x if not lines else 0.2, 0.2)
    x -= BLEED
    y = b["y"] - BLEED

    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(b["h"] + 0.35)
    )
    tf = box.text_frame
    # Every line break is the browser's own, so nothing may re-wrap here.
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    # Leave autofit off so Canva keeps the size we set rather than reflowing.
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for tag in ("a:normAutofit", "a:spAutoFit"):
        el = bodyPr.find(qn(tag))
        if el is not None:
            bodyPr.remove(el)

    def style(run, r):
        f = run.font
        f.name = r["family"]
        f.size = Pt(r["size"] * PX_TO_PT)
        f.bold = r["weight"] >= 600
        f.italic = r["italic"]
        colour, alpha = parse_color(r["color"])
        f.color.rgb = colour
        set_alpha(run, alpha)
        set_spacing(run, r.get("spacing", 0))

    text_runs = [r for r in b["runs"] if not r.get("br")]
    paras = []
    for i, ln in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paras.append(para)
        for part in ln["parts"]:
            r = text_runs[part["runIdx"]] if part["runIdx"] < len(text_runs) else text_runs[0]
            text = part["text"]
            if r.get("transform") == "uppercase":
                text = text.upper()
            if not text:
                continue
            run = para.add_run()
            run.text = text
            style(run, r)

    for para in paras:
        para.alignment = ALIGN.get(b["align"], PP_ALIGN.LEFT)
        para.line_spacing = Pt(b["lineHeightPx"] * PX_TO_PT)
        para.space_before = Pt(0)
        para.space_after = Pt(0)

    if HALF_LEADING:
        lead_in = (b["lineHeightPx"] - b["fontSizePx"]) / 2 / 96
        box.top = Emu(int(round((y - lead_in) * 914400)))
    return box


def crop_to_trim(png):
    """Cut the 0.125in bleed off a background render, once, beside the source."""
    from PIL import Image

    out = png.with_name(png.stem + "-trim.png")
    im = Image.open(png)
    dpi = im.width / (TRIM_W + 2 * BLEED)
    m = int(round(BLEED * dpi))
    im.crop((m, m, im.width - m, im.height - m)).save(out)
    return out


def main():
    layout = json.loads((SRC / "layout.json").read_text())

    prs = Presentation()
    prs.slide_width = Inches(TRIM_W)
    prs.slide_height = Inches(TRIM_H)
    blank = prs.slide_layouts[6]

    total = 0
    for name in ("outside", "inside"):
        data = layout[name]
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(crop_to_trim(SRC / f"bg-{name}.png")), 0, 0,
            width=prs.slide_width, height=prs.slide_height,
        )
        for b in data["blocks"]:
            add_block(slide, b)
        total += len(data["blocks"])
        print(f"  {name}: background + {len(data['blocks'])} text boxes")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"{OUT.relative_to(ROOT.parent)}  {OUT.stat().st_size / 1e6:.2f} MB  "
          f"2 slides, {total} editable text boxes")


if __name__ == "__main__":
    main()
